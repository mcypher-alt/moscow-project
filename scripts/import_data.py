"""Import source-obfuscated CSV/XLSX or GeoJSON through the validated REST API.

Canonical columns match docs/API.md. Original Russian reference/log headers are
also accepted. Credentials come from environment variables, never command args.
"""
import argparse
import csv
from datetime import datetime
import json
import math
import os
from pathlib import Path
import re
from urllib.request import Request, build_opener, HTTPCookieProcessor
from zoneinfo import ZoneInfo

ALIASES = {
    "ид_объект": "systemObjectId", "ид_канала_данных": "channelId", "ид_события": "id",
    "иерархия_уровень": "level", "вид_объекта": "objectKind", "родитель": "parentId",
    "диспетчерское_название_объекта": "dispatcherName", "тег_инженерной_системы": "systemTag",
    "название_датчика": "sensorName", "тип_инж_системы": "systemType", "тип_датчика": "sensorType",
    "тревожное": "isAlarm", "значение_датчика": "rawValue",
}


def rows(path):
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            iterator = workbook.active.iter_rows(values_only=True)
            headers = next(iterator)
            for row in iterator:
                yield {str(k): v for k, v in zip(headers, row) if k is not None and v is not None}
        finally:
            workbook.close()
    elif path.suffix.lower() in ('.json', '.geojson'):
        document = json.loads(path.read_text(encoding='utf-8-sig'))
        for feature in document['features']:
            geometry = feature['geometry']
            if geometry['type'] != 'Point':
                raise ValueError('Only WGS84 Point geometry is supported for object locations')
            lng, lat = geometry['coordinates'][:2]
            yield {**feature['properties'], 'longitude': lng, 'latitude': lat}
    else:
        with path.open(encoding='utf-8-sig', newline='') as stream:
            sample = stream.read(8192)
            stream.seek(0)
            yield from csv.DictReader(stream, dialect=csv.Sniffer().sniff(sample, delimiters=',;\t'))


def normalize(row, kind):
    result = {ALIASES.get(k, k): v for k, v in row.items() if v is not None and v != ''}
    if kind == 'objects': result['id'] = result.pop('systemObjectId', result.get('id'))
    if kind == 'channels': result['id'] = result.pop('channelId', result.get('id'))
    for key in ('id', 'systemObjectId', 'channelId', 'level', 'parentId'):
        if key in result: result[key] = str(result[key]) if key == 'id' and kind == 'telemetry' else int(result[key])
    for key in ('latitude', 'longitude', 'numericValue'):
        if key in result: result[key] = float(result[key])
    if 'wkt' in result:
        match = re.fullmatch(r'POINT\s*\(\s*([-+\d.eE]+)\s+([-+\d.eE]+)\s*\)', str(result.pop('wkt')), re.I)
        if not match: raise ValueError('Expected WGS84 POINT(longitude latitude)')
        result['longitude'], result['latitude'] = map(float, match.groups())
    if kind == 'telemetry':
        if 'recordedAt' not in result:
            result['recordedAt'] = f"{result.pop('дата')}T{result.pop('время')}"
        at = datetime.fromisoformat(str(result['recordedAt']))
        if at.tzinfo is None: at = at.replace(tzinfo=ZoneInfo('Europe/Moscow'))
        result['recordedAt'] = at.isoformat()
        alarm = str(result['isAlarm']).strip().lower()
        if alarm not in ('true', 'false', 't', 'f', '1', '0', 'да', 'нет', 'yes', 'no'):
            raise ValueError('Unknown alarm flag')
        result['isAlarm'] = alarm in ('true', 't', '1', 'да', 'yes')
        if 'rawValue' in result:
            result['rawValue'] = str(result['rawValue'])
            try:
                numeric = float(result['rawValue'].strip().replace(',', '.'))
                result['numericValue'] = numeric if math.isfinite(numeric) else None
            except ValueError: result['numericValue'] = None
    for key in ('startsAt', 'endsAt'):
        if key in result:
            at = datetime.fromisoformat(str(result[key]))
            if at.tzinfo is None: at = at.replace(tzinfo=ZoneInfo('Europe/Moscow'))
            result[key] = at.isoformat()
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('kind', choices=['objects', 'channels', 'telemetry', 'work-requests'])
    parser.add_argument('file', type=Path)
    parser.add_argument('--url', default='http://localhost:3000/api')
    parser.add_argument('--dry-run', action='store_true')
    parser.add_argument('--root-parent', type=int, help='External root ID omitted from the supplied object registry')
    args = parser.parse_args()
    opener = build_opener(HTTPCookieProcessor())

    def post(route, body):
        request = Request(args.url.rstrip('/') + route, json.dumps(body, allow_nan=False).encode(), {'Content-Type': 'application/json'})
        with opener.open(request, timeout=300) as response: return json.load(response)

    if not args.dry_run:
        post('/auth/login', {'email': os.environ['IMPORT_EMAIL'], 'password': os.environ['IMPORT_PASSWORD']})
    batch, total = [], 0
    source = rows(args.file)
    if args.kind == 'objects':
        source = sorted(source, key=lambda row: int(row.get('level', row.get('иерархия_уровень', 0))))
    for row in source:
        item = normalize(row, args.kind)
        if args.kind == 'objects' and item.get('parentId') == args.root_parent:
            item['parentId'] = None
        batch.append(item)
        if len(batch) == 500:
            if not args.dry_run: print(post('/integrations/' + args.kind, {'items': batch}))
            total += len(batch)
            batch = []
    if batch:
        if not args.dry_run: print(post('/integrations/' + args.kind, {'items': batch}))
        total += len(batch)
    print(f'{total} rows processed' + (' (dry run, no writes)' if args.dry_run else ''))


if __name__ == '__main__': main()
