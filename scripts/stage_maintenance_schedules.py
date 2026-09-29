"""Read organizer schedules into reviewable JSON; never infer operational object IDs.

Requires openpyxl. This is a staging export, not a database import or proof of work.
Merged cells resolve only to their actual anchor, never arbitrary forward filling.
"""
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/sources/chat-20260928'


def cell_value(sheet, row, column):
    cell = sheet.cell(row, column)
    for merged in sheet.merged_cells.ranges:
        if cell.coordinate in merged:
            cell = sheet.cell(merged.min_row, merged.min_col)
            break
    value = cell.value
    if isinstance(value, (datetime, date)):
        value = value.isoformat()
    return {'cell': cell.coordinate, 'value': value}


def main():
    records, files = [], []
    for path in sorted(SOURCE.glob('График*.xlsx')):
        book = load_workbook(path, data_only=False)
        files.append({'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
        sheet = book.active
        if 'ППР' in path.name:
            for row in range(10, 36):
                fields = {name: cell_value(sheet, row, column) for column, name in enumerate(
                    ['division', 'month', 'objectLabel', 'quantity', 'dismantling', 'delivery', 'collection', 'acceptance'], 1)}
                records.append({'kind': 'PPR', 'file': path.name, 'sheet': sheet.title, 'row': row,
                                'systemObjectId': None, 'fields': fields})
        else:
            group = None
            for row in range(6, sheet.max_row + 1):
                number = sheet.cell(row, 2).value
                if isinstance(number, (int, float)):
                    group = {'number': number, 'label': cell_value(sheet, row, 3)}
                equipment = sheet.cell(row, 4).value
                if not group or not equipment or equipment == 'Марка':
                    continue
                records.append({'kind': 'TO_TR', 'file': path.name, 'sheet': sheet.title, 'row': row,
                    'year': 2026, 'yearSource': 'G2 title and filename; sheet title says 2025',
                    'sourceGroup': group, 'systemObjectId': None,
                    'equipment': cell_value(sheet, row, 4), 'quantity': cell_value(sheet, row, 5),
                    'unit': cell_value(sheet, row, 6),
                    'months': [{'month': col - 6, **cell_value(sheet, row, col)} for col in range(7, 19)]})
        book.close()
    output = {'status': 'unmapped_reference_only',
              'warning': 'Anonymous source labels and row numbers are NOT operational object IDs. Monthly marks are not exact intervals or completed work.',
              'files': files, 'records': records}
    (SOURCE / 'maintenance-staging.json').write_text(json.dumps(output, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Staged {len(records)} source records; no database writes')


if __name__ == '__main__':
    main()
